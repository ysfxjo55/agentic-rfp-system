import os
import re
import unicodedata
from typing import List, Dict, Any, Optional
from pathlib import Path
import fitz  # PyMuPDF
import docx
from app.models.schemas import ExtractedBlock

# Arabic Unicode ranges (standard block, supplement, presentation forms A/B)
_ARABIC_CHAR_RE = re.compile(r'[؀-ۿݐ-ݿﭐ-﷿ﹰ-﻿]')


def _is_arabic_dominant(text: str) -> bool:
    """True if Arabic characters make up a significant share of non-whitespace text."""
    stripped = re.sub(r'\s+', '', text)
    if not stripped:
        return False
    arabic_count = len(_ARABIC_CHAR_RE.findall(stripped))
    return (arabic_count / len(stripped)) > 0.3


def _fix_rtl_line(line_text: str) -> str:
    """
    Fallback text-only RTL correction for non-PDF sources (DOCX/TXT), where
    no glyph coordinates are available. Reverses word order for Arabic-dominant
    lines and applies NFKC normalization to collapse Arabic Presentation Form
    ligatures (e.g. U+FEFB) back to standard letter sequences (e.g. "لا").
    PDF extraction uses coordinate-based word reordering instead (see
    `_reconstruct_pdf_line_from_words`), which is materially more reliable.
    """
    normalized = unicodedata.normalize('NFKC', line_text)
    if not _is_arabic_dominant(normalized):
        return normalized

    tokens = normalized.split()
    if len(tokens) <= 1:
        return normalized
    return " ".join(reversed(tokens))


def _reconstruct_pdf_line_from_words(words: List[tuple]) -> str:
    """
    Reconstructs a physical PDF line's logical reading order from individual
    word bounding boxes, rather than trusting the PDF content-stream order.
    Some PDF generators (including this document's source, a Saudi government
    tender portal) emit RTL Arabic word runs in visual (left-to-right screen
    position) order rather than logical (right-to-left reading) order; sorting
    by descending x-coordinate reconstructs the correct reading order directly
    from physical layout, which is robust regardless of how the source PDF's
    content stream ordered its glyphs.

    `words` is a list of (x0, y0, x1, y1, word, block_no, line_no, word_no) tuples
    for a single physical line, as returned by PyMuPDF's page.get_text("words").
    """
    combined = " ".join(w[4] for w in words)
    combined_nfkc = unicodedata.normalize('NFKC', combined)
    is_rtl = _is_arabic_dominant(combined_nfkc)
    ordered = sorted(words, key=lambda w: -w[0] if is_rtl else w[0])

    # Some Arabic glyphs (notably the lam-alef ligature "لا") get emitted by
    # this PDF's font as a separate word-tuple mid-word rather than joined
    # with its neighbors, splitting one logical word into 2-3 fragments (e.g.
    # "وسلامة" -> "وس" / "ﻼ" / "مة"). These fragments sit almost exactly
    # adjacent with near-zero physical gap, unlike genuine inter-word spacing
    # (empirically ~3.5+ points in this document vs <1 point for a same-word
    # split) — so join fragments directly (no space) when the gap is small,
    # and insert a normal space otherwise.
    GAP_THRESHOLD = 1.0
    parts: List[str] = []
    prev_edge: Optional[float] = None
    for w in ordered:
        x0, _, x1, *_rest = w[:4]
        text = w[4]
        if prev_edge is not None:
            gap = (prev_edge - x1) if is_rtl else (x0 - prev_edge)
            if gap >= GAP_THRESHOLD:
                parts.append(" ")
        parts.append(text)
        prev_edge = x0 if is_rtl else x1

    return unicodedata.normalize('NFKC', "".join(parts))


class DocumentParserService:
    @staticmethod
    def get_file_extension(file_path: str) -> str:
        return Path(file_path).suffix.lower()

    @classmethod
    def parse_document(cls, file_path: str) -> List[ExtractedBlock]:
        """
        Parses a PDF or DOCX document and returns a list of ExtractedBlocks
        annotated with page numbers, section titles, and block types.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Document not found at {file_path}")

        ext = cls.get_file_extension(file_path)
        if ext == ".pdf":
            return cls._parse_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return cls._parse_docx(file_path)
        elif ext in [".txt", ".md"]:
            return cls._parse_text(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}. Only PDF, DOCX, and TXT are supported.")

    @classmethod
    def get_page_count(cls, file_path: str) -> int:
        ext = cls.get_file_extension(file_path)
        if ext == ".pdf":
            try:
                doc = fitz.open(file_path)
                count = len(doc)
                doc.close()
                return count
            except Exception:
                return 1
        elif ext in [".docx", ".doc"]:
            # Approximation for DOCX (roughly 400 words per page)
            try:
                doc = docx.Document(file_path)
                words = sum(len(p.text.split()) for p in doc.paragraphs)
                return max(1, (words // 400) + 1)
            except Exception:
                return 1
        return 1

    @classmethod
    def _parse_pdf(cls, file_path: str) -> List[ExtractedBlock]:
        doc = fitz.open(file_path)
        blocks_out: List[ExtractedBlock] = []
        current_section = "Introduction & General Information"

        for page_num in range(len(doc)):
            page = doc[page_num]
            # page_num is 0-indexed in fitz, so 1-indexed for citations
            page_index = page_num + 1

            # Extract at word granularity (not "blocks"/"dict" text order) so each
            # physical line can be reconstructed from actual glyph x-coordinates.
            # Some PDF generators (including this document's source, a government
            # tender portal) emit RTL Arabic word runs in visual/content-stream
            # order rather than logical reading order; coordinate-based sorting
            # reconstructs correct reading order regardless of stream order.
            words = page.get_text("words")  # (x0, y0, x1, y1, word, block_no, line_no, word_no)

            blocks_map: "Dict[int, Dict[int, List[tuple]]]" = {}
            block_order: List[int] = []
            for w in words:
                block_no, line_no = w[5], w[6]
                if block_no not in blocks_map:
                    blocks_map[block_no] = {}
                    block_order.append(block_no)
                blocks_map[block_no].setdefault(line_no, []).append(w)

            for block_no in block_order:
                lines_map = blocks_map[block_no]
                corrected_lines: List[str] = []
                for line_no in sorted(lines_map.keys()):
                    line_words = lines_map[line_no]
                    line_text = _reconstruct_pdf_line_from_words(line_words).strip()
                    if line_text:
                        corrected_lines.append(line_text)

                if not corrected_lines:
                    continue

                text = "\n".join(corrected_lines)

                # Heading detection heuristic:
                # Short line (< 90 chars), doesn't end with a period, often capitalized or numbered
                first_line = corrected_lines[0]
                is_heading = False

                if len(first_line) < 90 and (
                    first_line.isupper()
                    or any(first_line.startswith(prefix) for prefix in ["Section", "Part", "Chapter", "1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9."])
                    or not first_line.endswith((".", ":", ";"))
                ) and len(corrected_lines) <= 2:
                    current_section = first_line
                    is_heading = True

                blocks_out.append(
                    ExtractedBlock(
                        text=text,
                        page_number=page_index,
                        section_title=current_section,
                        block_type="heading" if is_heading else "paragraph"
                    )
                )

        doc.close()
        return blocks_out

    @classmethod
    def _parse_docx(cls, file_path: str) -> List[ExtractedBlock]:
        doc = docx.Document(file_path)
        blocks_out: List[ExtractedBlock] = []
        current_section = "General Requirements"
        current_page = 1
        word_count = 0

        for p in doc.paragraphs:
            text = _fix_rtl_line(p.text.strip())
            if not text:
                continue

            word_count += len(text.split())
            current_page = max(1, (word_count // 400) + 1)

            style_name = p.style.name.lower() if p.style else ""
            is_heading = "heading" in style_name or "title" in style_name

            if is_heading:
                current_section = text

            blocks_out.append(
                ExtractedBlock(
                    text=text,
                    page_number=current_page,
                    section_title=current_section,
                    block_type="heading" if is_heading else "paragraph"
                )
            )

        # Also extract table contents
        for table in doc.tables:
            table_text = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_data:
                    table_text.append(" | ".join(row_data))
            if table_text:
                full_table = "\n".join(table_text)
                blocks_out.append(
                    ExtractedBlock(
                        text=full_table,
                        page_number=current_page,
                        section_title=current_section,
                        block_type="table"
                    )
                )

        return blocks_out

    @classmethod
    def _parse_text(cls, file_path: str) -> List[ExtractedBlock]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            lines = [_fix_rtl_line(line.strip()) for line in f.readlines()]

        blocks_out: List[ExtractedBlock] = []
        current_section = "General Information"
        current_page = 1
        line_count = 0

        for line in lines:
            if not line:
                continue
            line_count += 1
            current_page = (line_count // 15) + 1

            if (
                line.startswith("#")
                or line.upper().startswith("SECTION ")
                or (len(line) < 90 and line.isupper() and not line.endswith("."))
            ):
                current_section = line.lstrip("#").strip()
                blocks_out.append(
                    ExtractedBlock(
                        text=line,
                        page_number=current_page,
                        section_title=current_section,
                        block_type="heading"
                    )
                )
            else:
                blocks_out.append(
                    ExtractedBlock(
                        text=line,
                        page_number=current_page,
                        section_title=current_section,
                        block_type="paragraph"
                    )
                )
        return blocks_out
